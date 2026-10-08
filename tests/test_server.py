import base64
import http.client
import io
import json
import threading
import unittest
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from tests.helpers import TEST_SEED, fresh_db
from accsoft import licensing
from accsoft.server import create_server


class Client:
    def __init__(self, port):
        self.port, self.cookie, self.csrf = port, None, None

    def call(self, method, path, body=None, headers=None, host=None, csrf=True):
        c = http.client.HTTPConnection("127.0.0.1", self.port)
        h = {"Host": host or f"127.0.0.1:{self.port}", **(headers or {})}
        if self.cookie:
            h["Cookie"] = self.cookie
        if self.csrf and csrf:
            h["X-CSRF-Token"] = self.csrf
        data = json.dumps(body) if body is not None else None
        c.request(method, path, data, h)
        r = c.getresponse()
        raw = r.read()
        sc = r.getheader("Set-Cookie")
        if sc:
            self.cookie = sc.split(";")[0]
        try:
            return r.status, json.loads(raw), r
        except ValueError:
            return r.status, raw, r

    def login(self, user="admin", pw="password123"):
        st, _, _ = self.call("POST", "/api/setup", {"username": user, "password": pw, "shop_name": "فروشگاه تست"})
        st, d, _ = self.call("GET", "/api/state")
        self.csrf = d["csrf"]


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.db = fresh_db()
        self.srv = create_server(self.db)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.c = Client(self.srv.app.port)
        self.c.login()

    def tearDown(self):
        self.srv.shutdown()
        self.srv.server_close()

    def test_auth_required_and_setup_once(self):
        anon = Client(self.srv.app.port)
        self.assertEqual(anon.call("GET", "/api/products")[0], 401)
        self.assertEqual(anon.call("GET", "/export/products.xlsx")[0], 401)
        self.assertEqual(anon.call("GET", "/print/products")[0], 401)
        self.assertEqual(anon.call("POST", "/api/setup", {"username": "evil", "password": "12345678"})[0], 400)
        self.assertEqual(anon.call("POST", "/api/login", {"username": "admin", "password": "wrong"})[0], 400)

    def test_login_lockout(self):
        anon = Client(self.srv.app.port)
        for _ in range(5):
            anon.call("POST", "/api/login", {"username": "admin", "password": "x"})
        st, d, _ = anon.call("POST", "/api/login", {"username": "admin", "password": "password123"})
        self.assertEqual(st, 400)
        self.assertIn("صبر", d["error"])

    def test_csrf_host_and_origin(self):
        self.assertEqual(self.c.call("POST", "/api/expenses", {"title": "a", "amount": 1}, csrf=False)[0], 403)
        self.assertEqual(self.c.call("GET", "/api/products", host="evil.com")[0], 403)  # DNS rebinding
        self.assertEqual(self.c.call("POST", "/api/expenses", {"title": "a", "amount": 1},
                                     headers={"Origin": "http://evil.com"})[0], 403)
        self.assertEqual(self.c.call("POST", "/api/expenses", {"title": "a", "amount": 1})[0], 200)

    def test_security_headers_and_static_traversal(self):
        st, _, r = self.c.call("GET", "/")
        self.assertEqual(st, 200)
        self.assertIn("default-src 'self'", r.getheader("Content-Security-Policy"))
        self.assertEqual(r.getheader("X-Content-Type-Options"), "nosniff")
        for bad in ("/..%2f..%2fetc/passwd", "/../server.py", "/%2e%2e/db.py", "/a/b"):
            self.assertEqual(self.c.call("GET", bad)[0], 404, bad)
        self.assertEqual(self.c.call("GET", "/img/../../x.png")[0], 404)

    def test_sale_flow_invoice_report_exports(self):
        _, d, _ = self.c.call("POST", "/api/products", {"name": "<img src=x onerror=1>", "price": 1000, "cost": 400, "stock": 5})
        pid = d["id"]
        st, s, _ = self.c.call("POST", "/api/sales", {"items": [{"product_id": pid, "qty": 2}], "discount": 100,
                                                      "customer": {"phone": "09120000000", "first_name": "ع"}})
        self.assertEqual(st, 200)
        st, page, _ = self.c.call("GET", f"/print/invoice/{s['id']}")
        self.assertNotIn(b"<img src=x", page)  # XSS در فاکتور escape شده
        self.assertIn(b"&lt;img", page)
        st, rep, _ = self.c.call("GET", "/api/report?kind=day")
        self.assertEqual((rep["totals"]["revenue"], rep["totals"]["net_profit"]), (1900, 1100))
        st, x, r = self.c.call("GET", "/export/products.xlsx")
        ws = load_workbook(io.BytesIO(x)).active
        self.assertTrue(str(ws["B2"].value).startswith("<img"))
        for path in ("/export/sales.xlsx", "/export/customers.xlsx", "/export/report.xlsx?kind=month", "/print/report?kind=week"):
            self.assertEqual(self.c.call("GET", path)[0], 200, path)

    def test_excel_formula_injection_neutralised(self):
        self.c.call("POST", "/api/products", {"name": "=HYPERLINK(\"http://evil\")", "price": 1})
        _, x, _ = self.c.call("GET", "/export/products.xlsx")
        self.assertTrue(load_workbook(io.BytesIO(x)).active["B2"].value.startswith("'="))

    def test_image_upload_rejects_non_image(self):
        _, d, _ = self.c.call("POST", "/api/products", {"name": "p", "price": 1})
        bad = base64.b64encode(b"<script>alert(1)</script>").decode()
        self.assertEqual(self.c.call("POST", f"/api/products/{d['id']}/image", {"data": bad})[0], 400)
        good = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"0" * 20).decode()
        st, r, _ = self.c.call("POST", f"/api/products/{d['id']}/image", {"data": good})
        self.assertEqual(st, 200)
        st, img, resp = self.c.call("GET", f"/img/{r['image']}")
        self.assertEqual((st, resp.getheader("Content-Type")), (200, "image/png"))

    def test_secrets_never_returned(self):
        self.c.call("PUT", "/api/settings", {"woo_key": "ck_SECRET123", "sms_apikey": "smsSECRET", "shop_name": "x"})
        _, s, _ = self.c.call("GET", "/api/settings")
        self.assertEqual(s["woo_key"], "********")
        self.assertNotIn("SECRET", json.dumps(s))
        self.assertNotIn("ck_SECRET123", self.db.get("woo_key"))  # روی دیسک رمز شده
        self.c.call("PUT", "/api/settings", {"woo_key": "********"})  # مقدار ماسک‌شده کلید را بازنویسی نکند
        from accsoft import security
        self.assertEqual(security.unseal(self.db.get("woo_key")), "ck_SECRET123")

    def test_license_gate_blocks_writes_when_expired_and_unlocks(self):
        self.db.set("first_run", (datetime.now() - timedelta(days=30)).isoformat())
        st, d, _ = self.c.call("POST", "/api/expenses", {"title": "a", "amount": 1})
        self.assertEqual(st, 402)
        self.assertEqual(self.c.call("GET", "/api/products")[0], 200)  # مشاهده آزاد
        self.assertEqual(self.c.call("GET", "/export/products.xlsx")[0], 200)
        self.assertEqual(self.c.call("POST", "/api/license", {"key": "junk"})[0], 400)
        key = licensing.make_license(TEST_SEED, "12m", "تست")
        self.assertEqual(self.c.call("POST", "/api/license", {"key": key})[0], 200)
        self.assertEqual(self.c.call("POST", "/api/expenses", {"title": "a", "amount": 1})[0], 200)

    def test_errors_do_not_leak_internals(self):
        st, d, _ = self.c.call("POST", "/api/sales", {"items": [{"product_id": 999, "qty": 1}]})
        self.assertEqual(st, 400)
        st, d, _ = self.c.call("PUT", "/api/settings", {"usd_rate": "abc"})
        self.assertEqual(st, 400)
        self.assertEqual(self.c.call("POST", "/api/sales", [1, 2])[0], 400)

    def test_sql_injection_attempts_are_inert(self):
        self.c.call("POST", "/api/products", {"name": "x", "price": 1})
        st, d, _ = self.c.call("GET", "/api/products?q=%27%3B%20DROP%20TABLE%20products%3B--")
        self.assertEqual((st, d), (200, []))
        self.assertEqual(len(self.c.call("GET", "/api/products")[1]), 1)


if __name__ == "__main__":
    unittest.main()


class ResetTests(unittest.TestCase):
    def test_reset_flow(self):
        import os
        import tempfile
        os.environ["ACCSOFT_DATA"] = tempfile.mkdtemp()
        self.addCleanup(os.environ.pop, "ACCSOFT_DATA", None)
        db = fresh_db()
        srv = create_server(db)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        self.addCleanup(srv.server_close)
        self.addCleanup(srv.shutdown)
        c = Client(srv.app.port)
        c.login("boss", "Abcdef123")
        c.call("POST", "/api/logout", {})
        c = Client(srv.app.port)
        st, r, _ = c.call("POST", "/api/reset/request", {})
        self.assertEqual(st, 200)
        code = open(r["path"], encoding="utf-8").read().split(":")[1].split()[0]
        st, d, _ = c.call("POST", "/api/reset/confirm", {"code": "ZZZZZZZZ", "password": "NewPass123"})
        self.assertEqual(st, 400)
        st, d, _ = c.call("POST", "/api/reset/confirm", {"code": code, "password": "short"})
        self.assertEqual(st, 400)
        st, d, _ = c.call("POST", "/api/reset/confirm", {"code": code, "password": "NewPass123"})
        self.assertEqual((st, d["username"]), (200, "boss"))
        self.assertFalse(os.path.exists(r["path"]))
        self.assertEqual(c.call("POST", "/api/login", {"username": "BOSS", "password": "NewPass123"})[0], 200)   # حروف بزرگ/کوچک مهم نیست
        c2 = Client(srv.app.port)
        self.assertEqual(c2.call("POST", "/api/login", {"username": "boss", "password": "Abcdef123"})[0], 400)
        self.assertEqual(c2.call("POST", "/api/reset/confirm", {"code": code, "password": "Another123"})[0], 400)   # یک‌بار مصرف


class AccountTests(unittest.TestCase):
    """ورود/ثبت‌نام با حساب وبیکری؛ سرور وبیکری با monkeypatch شبیه‌سازی می‌شود."""

    def setUp(self):
        from unittest import mock
        from accsoft import account
        self.site_pw = {"u@x.ir": "SitePass123"}   # «پایگاه کاربران» سایت

        def fake_post(path, body):
            from accsoft.services import AppError
            if path == "/register":
                if body["email"] in self.site_pw:
                    raise AppError("با این ایمیل قبلاً حساب ساخته شده")
                self.site_pw[body["email"]] = body["password"]
                return {"ok": True, "email": body["email"], "name": body["name"], "licenses": []}
            if self.site_pw.get(body["login"].lower()) != body["password"]:
                raise AppError("ایمیل/نام کاربری یا رمز نادرست است")
            return {"ok": True, "email": body["login"].lower(), "name": "n", "licenses": []}
        self.patches = [mock.patch.object(account, "REQUIRED", True), mock.patch.object(account, "_post", fake_post)]
        for p in self.patches:
            p.start()
        self.db = fresh_db()
        self.srv = create_server(self.db)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.port = self.srv.app.port

    def tearDown(self):
        for p in self.patches:
            p.stop()
        self.srv.shutdown()
        self.srv.server_close()

    def test_setup_requires_site_account_and_password_syncs_from_site(self):
        c = Client(self.port)
        self.assertTrue(c.call("GET", "/api/state")[1]["account_required"])
        self.assertEqual(c.call("POST", "/api/setup", {"mode": "login", "email": "u@x.ir", "password": "wrong-pass1", "shop_name": "s"})[0], 400)
        st, _, _ = c.call("POST", "/api/setup", {"mode": "login", "email": "U@x.ir", "password": "SitePass123", "shop_name": "s"})
        self.assertEqual(st, 200)
        u = self.db.one("SELECT username,account FROM users")
        self.assertEqual((u["username"], u["account"]), ("u@x.ir", "u@x.ir"))
        # رمز در سایت عوض می‌شود → ورود با رمز جدید (تأیید آنلاین) و رمز محلی همگام می‌شود
        c.call("POST", "/api/logout", {})
        self.site_pw["u@x.ir"] = "BrandNew456"
        c2 = Client(self.port)
        self.assertEqual(c2.call("POST", "/api/login", {"username": "u@x.ir", "password": "BrandNew456"})[0], 200)
        self.assertEqual(Client(self.port).call("POST", "/api/login", {"username": "u@x.ir", "password": "SitePass123"})[0], 400)   # رمز قدیمی دیگر کار نمی‌کند

    def test_register_and_link(self):
        c = Client(self.port)
        st, _, _ = c.call("POST", "/api/setup", {"mode": "register", "email": "new@x.ir", "password": "Abcdef123", "name": "علی", "shop_name": "s"})
        self.assertEqual(st, 200)
        self.assertIn("new@x.ir", self.site_pw)
        dup = Client(self.port).call("POST", "/api/setup", {"mode": "register", "email": "u@x.ir", "password": "Abcdef123"})
        self.assertEqual(dup[0], 400)   # قبلاً حساب دارد / نصب انجام شده
        c.csrf = c.call("GET", "/api/state")[1]["csrf"]
        st, d, _ = c.call("POST", "/api/account/link", {"email": "u@x.ir", "password": "SitePass123"})
        self.assertEqual((st, d["account"]), (200, "u@x.ir"))
